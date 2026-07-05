import json
import os
import sqlite3
from contextlib import asynccontextmanager, contextmanager
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Optional

import httpx
from fastapi import FastAPI, HTTPException, Query
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

BASE_DIR = Path(__file__).resolve().parent
DB_PATH = Path(os.environ.get("WEBAPP_DB", BASE_DIR / "data" / "records.db"))
UROTRIAGE_API_URL = os.environ.get("UROTRIAGE_API_URL", "http://127.0.0.1:8000").rstrip("/")
AUTO_EXPORT_PATH = Path(os.environ.get("WEBAPP_AUTO_XLSX", DB_PATH.parent / "анкета_актуальная.xlsx"))
STATIC_DIR = BASE_DIR / "static"

COMPLAINT_PHRASES = {
    "no_voiding": "невозможность самостоятельного мочеиспускания",
    "suprapubic_pain": "боли над лоном",
    "severe_pain": "выраженная боль",
    "urge": "императивные позывы к мочеиспусканию",
    "fever_chills": "лихорадка, озноб",
    "gross_hematuria": "макрогематурия",
    "blood_clots": "макрогематурия со сгустками",
    "weakness": "слабость, головокружение",
}

HYDRONEPHROSIS_PHRASES = {
    "none": "гидронефроза нет",
    "unilateral": "односторонний гидронефроз",
    "bilateral": "двусторонний гидронефроз",
    "pyeloectasia": "пиелоэктазия",
}

DRAINAGE_LABELS = {
    "urethral": "уретральный катетер",
    "cystostomy": "цистостомия",
    "trocar": "троакарная эпицистостомия",
    "none": "дренирование не потребовалось",
}

ATTEMPT_LABELS = {
    "first_try": "катетеризация с первой попытки",
    "difficult": "катетеризация с техническими трудностями",
    "failed": "безуспешная попытка катетеризации",
}

ONSET_LABELS = {
    "lt6": "менее 6 часов",
    "6to24": "6–24 часа",
    "gt24": "более 24 часов",
    "unknown": "не установлено",
}

EXPORT_COLUMNS = [
    ("case_id", "№ истории болезни"),
    ("created_at", "Дата записи"),
    ("updated_at", "Дата изменения"),
    ("doctor", "Врач"),
    ("zone", "Зона светофора"),
    ("admitted_at", "Дата и время поступления"),
    ("age", "Возраст"),
    ("onset_category", "Время от начала задержки"),
    ("complaints_flags", "Жалобы (структурно)"),
    ("complaints_text", "Жалобы (текст)"),
    ("diagnosis_text", "Диагноз при поступлении"),
    ("disease_duration_years", "Длительность ДГПЖ, лет"),
    ("prior_ozm_count", "ОЗМ ранее, раз"),
    ("prior_residual", "Остаточная моча ранее"),
    ("bladder_stones", "Камни мочевого пузыря"),
    ("prior_prostate_surgery", "Операции на простате ранее"),
    ("alpha_blockers", "Альфа-адреноблокаторы"),
    ("reductase_inhibitors", "Ингибиторы 5-АР"),
    ("psa", "ПСА, нг/мл"),
    ("temp", "Температура, °C"),
    ("bp_sys", "АД систолическое"),
    ("bp_dia", "АД диастолическое"),
    ("pulse", "Пульс"),
    ("wbc", "Лейкоциты, ×10⁹/л"),
    ("hemoglobin", "Гемоглобин, г/л"),
    ("creatinine", "Креатинин, мкмоль/л"),
    ("urea", "Мочевина, ммоль/л"),
    ("crp", "СРБ, мг/л"),
    ("pct", "Прокальцитонин, нг/мл"),
    ("prostate_vol", "Объём простаты, см³"),
    ("median_lobe", "Средняя доля"),
    ("median_lobe_size", "Размер средней доли, мм"),
    ("residual_urine_ml", "Остаточная моча / эвакуировано, мл"),
    ("hydronephrosis", "Гидронефроз"),
    ("imaging_text", "УЗИ/КТ заключение"),
    ("drainage_method", "Метод дренирования"),
    ("catheter_attempt", "Катетеризация"),
    ("had_operation", "Операция за госпитализацию"),
    ("operation_txt", "Вид операции"),
    ("hours_to_oper", "Часов до операции"),
    ("outcome_txt", "Исход госпитализации"),
    ("complications_txt", "Осложнения"),
    ("readmission", "Повторная госпитализация"),
]

BOOL_LABELS = {True: "Да", False: "Нет", None: ""}

OUTCOME_KEYS = ["had_operation", "operation_txt", "hours_to_oper", "outcome_txt", "complications_txt", "readmission"]


class RecordIn(BaseModel):
    payload: dict = Field(...)


class TriagePreviewIn(BaseModel):
    payload: dict = Field(...)


def now_iso():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


@contextmanager
def db():
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    try:
        yield conn
        conn.commit()
    finally:
        conn.close()


def init_db():
    with db() as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS records (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                case_id TEXT NOT NULL,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                doctor TEXT,
                payload TEXT NOT NULL,
                triage TEXT,
                zone TEXT,
                outcome_filled INTEGER NOT NULL DEFAULT 0
            )
            """
        )
        conn.execute("CREATE INDEX IF NOT EXISTS idx_records_case ON records(case_id)")


def _num(payload, key):
    v = payload.get(key)
    if v is None or v == "":
        return None
    try:
        return float(str(v).replace(",", "."))
    except ValueError:
        return None


def build_predict_input(payload):
    complaint_parts = [COMPLAINT_PHRASES[c] for c in payload.get("complaints_flags") or [] if c in COMPLAINT_PHRASES]
    free = (payload.get("complaints_text") or "").strip()
    if free:
        complaint_parts.append(free)
    if payload.get("catheter_attempt") == "failed":
        complaint_parts.append(ATTEMPT_LABELS["failed"])

    imaging_parts = []
    hydro = payload.get("hydronephrosis")
    if hydro in HYDRONEPHROSIS_PHRASES:
        imaging_parts.append(HYDRONEPHROSIS_PHRASES[hydro])
    residual = _num(payload, "residual_urine_ml")
    if residual is not None:
        imaging_parts.append(f"остаточная моча {residual:.0f} мл")
    prostate = _num(payload, "prostate_vol")
    if prostate is not None:
        imaging_parts.append(f"объём предстательной железы {prostate:.0f} см3")
    if payload.get("median_lobe") is True:
        size = _num(payload, "median_lobe_size")
        imaging_parts.append("средняя доля" + (f" {size:.0f} мм" if size is not None else ""))
    imaging_free = (payload.get("imaging_text") or "").strip()
    if imaging_free:
        imaging_parts.append(imaging_free)

    return {
        "age": _num(payload, "age"),
        "sex": "Мужской",
        "temp": _num(payload, "temp"),
        "bp_sys": _num(payload, "bp_sys"),
        "bp_dia": _num(payload, "bp_dia"),
        "pulse": _num(payload, "pulse"),
        "creatinine": _num(payload, "creatinine"),
        "urea": _num(payload, "urea"),
        "crp": _num(payload, "crp"),
        "hemoglobin": _num(payload, "hemoglobin"),
        "wbc": _num(payload, "wbc"),
        "residual_urine_ml": residual,
        "diagnosis_text": (payload.get("diagnosis_text") or "").strip() or None,
        "complaints_text": ". ".join(complaint_parts) or None,
        "imaging_text": ". ".join(imaging_parts) or None,
    }


async def call_triage(payload):
    try:
        async with httpx.AsyncClient(timeout=6.0) as client:
            resp = await client.post(f"{UROTRIAGE_API_URL}/predict", json=build_predict_input(payload))
            resp.raise_for_status()
            return resp.json(), None
    except Exception as exc:
        return None, str(exc)


def outcome_is_filled(payload):
    return any(payload.get(k) not in (None, "", []) for k in OUTCOME_KEYS)


def row_to_summary(row):
    triage = json.loads(row["triage"]) if row["triage"] else None
    return {
        "id": row["id"],
        "case_id": row["case_id"],
        "created_at": row["created_at"],
        "updated_at": row["updated_at"],
        "doctor": row["doctor"],
        "zone": row["zone"],
        "outcome_filled": bool(row["outcome_filled"]),
        "fired_count": len(triage["rule"]["fired_criteria"]) if triage else None,
    }


def row_to_full(row):
    return {
        **row_to_summary(row),
        "payload": json.loads(row["payload"]),
        "triage": json.loads(row["triage"]) if row["triage"] else None,
    }


@asynccontextmanager
async def lifespan(_app):
    init_db()
    yield


app = FastAPI(title="UroTriage Webapp", version="1.0.0", lifespan=lifespan)


@app.post("/api/records")
async def create_record(body: RecordIn):
    payload = body.payload
    case_id = (payload.get("case_id") or "").strip()
    if not case_id:
        raise HTTPException(status_code=422, detail="Не указан № истории болезни")
    triage, triage_error = await call_triage(payload)
    now = now_iso()
    with db() as conn:
        cur = conn.execute(
            "INSERT INTO records (case_id, created_at, updated_at, doctor, payload, triage, zone, outcome_filled) VALUES (?,?,?,?,?,?,?,?)",
            (
                case_id,
                now,
                now,
                (payload.get("doctor") or "").strip() or None,
                json.dumps(payload, ensure_ascii=False),
                json.dumps(triage, ensure_ascii=False) if triage else None,
                triage["final_zone"] if triage else None,
                int(outcome_is_filled(payload)),
            ),
        )
        record_id = cur.lastrowid
    save_auto_export()
    return {"id": record_id, "triage": triage, "triage_error": triage_error}


@app.put("/api/records/{record_id}")
async def update_record(record_id: int, body: RecordIn):
    with db() as conn:
        row = conn.execute("SELECT * FROM records WHERE id = ?", (record_id,)).fetchone()
    if row is None:
        raise HTTPException(status_code=404, detail="Запись не найдена")
    payload = {**json.loads(row["payload"]), **body.payload}
    case_id = (payload.get("case_id") or "").strip()
    if not case_id:
        raise HTTPException(status_code=422, detail="Не указан № истории болезни")
    triage, triage_error = await call_triage(payload)
    with db() as conn:
        conn.execute(
            "UPDATE records SET case_id=?, updated_at=?, doctor=?, payload=?, triage=?, zone=?, outcome_filled=? WHERE id=?",
            (
                case_id,
                now_iso(),
                (payload.get("doctor") or "").strip() or None,
                json.dumps(payload, ensure_ascii=False),
                json.dumps(triage, ensure_ascii=False) if triage else row["triage"],
                triage["final_zone"] if triage else row["zone"],
                int(outcome_is_filled(payload)),
                record_id,
            ),
        )
    save_auto_export()
    return {"id": record_id, "triage": triage, "triage_error": triage_error}


@app.get("/api/records")
def list_records(
    q: Optional[str] = Query(None),
    zone: Optional[str] = Query(None),
    missing_outcome: bool = Query(False),
    limit: int = Query(200, ge=1, le=1000),
):
    sql = "SELECT * FROM records"
    conds, args = [], []
    if q:
        conds.append("case_id LIKE ?")
        args.append(f"%{q.strip()}%")
    if zone:
        conds.append("zone = ?")
        args.append(zone)
    if missing_outcome:
        conds.append("outcome_filled = 0")
    if conds:
        sql += " WHERE " + " AND ".join(conds)
    sql += " ORDER BY id DESC LIMIT ?"
    args.append(limit)
    with db() as conn:
        rows = conn.execute(sql, args).fetchall()
    return {"records": [row_to_summary(r) for r in rows]}


@app.get("/api/records/{record_id}")
def get_record(record_id: int):
    with db() as conn:
        row = conn.execute("SELECT * FROM records WHERE id = ?", (record_id,)).fetchone()
    if row is None:
        raise HTTPException(status_code=404, detail="Запись не найдена")
    return row_to_full(row)


@app.post("/api/triage")
async def triage_preview(body: TriagePreviewIn):
    triage, triage_error = await call_triage(body.payload)
    return {"triage": triage, "triage_error": triage_error}


@app.get("/api/stats")
def stats():
    with db() as conn:
        rows = conn.execute("SELECT zone, outcome_filled, created_at FROM records").fetchall()
    total = len(rows)
    by_zone = {"red": 0, "yellow": 0, "green": 0, "unknown": 0}
    missing_outcome = 0
    last7 = 0
    week_ago = datetime.now(timezone.utc).timestamp() - 7 * 86400
    for r in rows:
        by_zone[r["zone"] if r["zone"] in by_zone else "unknown"] += 1
        if not r["outcome_filled"]:
            missing_outcome += 1
        try:
            if datetime.fromisoformat(r["created_at"]).timestamp() >= week_ago:
                last7 += 1
        except ValueError:
            pass
    return {"total": total, "last7": last7, "by_zone": by_zone, "missing_outcome": missing_outcome}


def build_workbook():
    from openpyxl import Workbook
    from openpyxl.styles import Font

    with db() as conn:
        rows = conn.execute("SELECT * FROM records ORDER BY id").fetchall()
    wb = Workbook()
    ws = wb.active
    ws.title = "Анкета"
    ws.append([title for _, title in EXPORT_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        payload = json.loads(row["payload"])
        merged = {**payload, "case_id": row["case_id"], "created_at": row["created_at"], "updated_at": row["updated_at"], "doctor": row["doctor"], "zone": row["zone"]}
        line = []
        for key, _title in EXPORT_COLUMNS:
            v = merged.get(key)
            if isinstance(v, bool) or v is None:
                v = BOOL_LABELS.get(v, v)
            elif key == "complaints_flags":
                v = "; ".join(COMPLAINT_PHRASES.get(c, c) for c in v) if isinstance(v, list) else v
            elif key == "hydronephrosis":
                v = HYDRONEPHROSIS_PHRASES.get(v, v)
            elif key == "drainage_method":
                v = DRAINAGE_LABELS.get(v, v)
            elif key == "catheter_attempt":
                v = ATTEMPT_LABELS.get(v, v)
            elif key == "onset_category":
                v = ONSET_LABELS.get(v, v)
            line.append(v)
        ws.append(line)
    for idx, _ in enumerate(EXPORT_COLUMNS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = 22
    return wb


def save_auto_export():
    try:
        wb = build_workbook()
        tmp = AUTO_EXPORT_PATH.with_suffix(".tmp")
        wb.save(tmp)
        tmp.replace(AUTO_EXPORT_PATH)
    except Exception:
        pass


@app.get("/api/export.xlsx")
def export_xlsx():
    wb = build_workbook()
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"urotriage_records_{datetime.now().strftime('%Y-%m-%d')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/api/health")
async def health():
    triage_ok = False
    triage_info = None
    try:
        async with httpx.AsyncClient(timeout=3.0) as client:
            resp = await client.get(f"{UROTRIAGE_API_URL}/health")
            resp.raise_for_status()
            triage_info = resp.json()
            triage_ok = True
    except Exception:
        pass
    with db() as conn:
        total = conn.execute("SELECT COUNT(*) FROM records").fetchone()[0]
    return {"status": "ok", "records": total, "triage_api": {"reachable": triage_ok, "url": UROTRIAGE_API_URL, "info": triage_info}}


@app.get("/")
def index():
    return FileResponse(STATIC_DIR / "index.html")


app.mount("/", StaticFiles(directory=STATIC_DIR), name="static")
